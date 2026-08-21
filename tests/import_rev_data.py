import openpyxl
import json
import paramiko

# ── 1. Read rev/ xlsx files locally ──────────────────────────────
item_wb = openpyxl.load_workbook('rev/Report Item Sales - 01-07-2026 - 31-07-2026 - Sekar Pizza.xlsx', read_only=True)
item_ws = item_wb['Sheet1']
item_rows = list(item_ws.iter_rows(values_only=True))

pay_wb = openpyxl.load_workbook('rev/Report Payment Methods - 01-07-2026 - 31-07-2026 - Sekar Pizza.xlsx', read_only=True)
pay_ws = pay_wb['Sheet1']
pay_rows = list(pay_ws.iter_rows(values_only=True))

sum_wb = openpyxl.load_workbook('rev/Report Sales Summary - 01-07-2026 - 31-07-2026 - Sekar Pizza.xlsx', read_only=True)
sum_ws = sum_wb['Sales Summary']
sum_rows = list(sum_ws.iter_rows(values_only=True))

# ── 2. Parse ──────────────────────────────────────────────────────
# Sales summary
summary = {}
for r in sum_rows:
    if r[0] and r[0] != 'Type':
        summary[r[0]] = r[1]

gross = summary.get('Gross Sales', 0)
discount = abs(summary.get('Discount', 0))
net = summary.get('Net Sales', 0)
gratuity = summary.get('Gratuity', 0)
total_collected = summary.get('Total Collected', 0)

# Payment methods
payments = {}
for r in pay_rows[1:]:
    if r[0]:
        payments[r[0]] = {'tx': int(r[1] or 0), 'amount': int(r[2] or 0)}

settle_cash = payments.get('Cash - Cash', {}).get('amount', 0)
settle_qris = payments.get('Digital Payments - Gopay', {}).get('amount', 0)
tx_count = sum(p['tx'] for p in payments.values())

# Item sales
items = []
for r in item_rows[1:]:
    if not r[0]:
        continue
    name = r[0]
    variant = r[1] or ''
    category = r[2] or ''
    sku = r[3] or ''
    qty = int(r[4] or 0)
    gross_sales = int(r[6] or 0)
    disc = int(r[7] or 0)
    refund = int(r[8] or 0)
    net_sales = int(r[9] or 0)
    items.append({
        'item_name': f'{name} {variant}'.strip(),
        'sku': sku,
        'category': category,
        'qty': str(qty),
        'gross_sales': str(gross_sales),
        'discount': str(disc),
        'refund': str(refund),
        'net_sales': str(net_sales),
    })

print(f'items: {len(items)}')
print(f'gross={gross} net={net} discount={discount} gratuity={gratuity} total={total_collected}')
print(f'cash={settle_cash} qris={settle_qris} tx={tx_count}')

# ── 3. Build payload ──────────────────────────────────────────────
BRAND_ID = 'BR-002'
BRAND_NAME = 'Sekarpizza'
OUTLET_ID = 'OL-003'
OUTLET_NAME = 'Sekarpizza Demangan'
DATE = '2026-07-31'

payload = {
    'pos_daily': [{
        'pos_id': 'POS-SKP-20260731',
        'date': DATE,
        'brand_id': BRAND_ID,
        'brand_name': BRAND_NAME,
        'outlet_id': OUTLET_ID,
        'outlet_name': OUTLET_NAME,
        'gross_sales': str(int(gross)),
        'net_sales': str(int(net)),
        'discount': str(int(discount)),
        'refund': '0',
        'void': '0',
        'tax': '0',
        'service_charge': str(int(gratuity)),
        'settle_cash': str(int(settle_cash)),
        'settle_qris': str(int(settle_qris)),
        'settle_card': '0',
        'settle_transfer': '0',
        'settle_marketplace': '0',
        'total_settlement': str(int(total_collected)),
        'settlement_difference': str(int(total_collected) - int(settle_cash) - int(settle_qris)),
        'transaction_count': str(tx_count),
        'aov': str(int(net / tx_count)) if tx_count else '0',
        'cashier': '',
        'shift': '',
        'payment_method': '',
        'source': 'moka_import',
        'source_ref': 'rev/Report Sales Summary - Sekar Pizza',
        'notes': 'Import bulan Juli 2026 dari report Moka Sekar Pizza',
        'source_module': 'finance',
        'source_transaction_id': '',
        'payment_source': '',
        'linked_expense_id': '',
        'linked_supplier_invoice_id': '',
        'linked_petty_cash_id': '',
        'created_by': 'U-001',
        'created_at': '2026-08-19 17:00:00',
        'updated_at': '2026-08-19 17:00:00',
    }],
    'pos_items': [
        {
            'pos_item_id': f'POSI-SKP-{i+1:04d}',
            'date': DATE,
            'brand_id': BRAND_ID,
            'brand_name': BRAND_NAME,
            'outlet_id': OUTLET_ID,
            'outlet_name': OUTLET_NAME,
            'item_name': it['item_name'],
            'sku': it['sku'],
            'category': it['category'],
            'qty': it['qty'],
            'gross_sales': it['gross_sales'],
            'discount': it['discount'],
            'refund': it['refund'],
            'net_sales': it['net_sales'],
            'source': 'moka_import',
            'created_at': '2026-08-19 17:00:00',
        }
        for i, it in enumerate(items)
    ],
}

# ── 4. Upload to VPS + run import ─────────────────────────────────
ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect('187.77.114.168', username='dev', password='password')

sftp = ssh.open_sftp()
with sftp.file('/home/dev/ykp/ykp-finance-v1/import_rev.json', 'w') as f:
    f.write(json.dumps(payload))
sftp.close()

NODE = r"""
const { google } = require('googleapis');
const fs = require('fs');
const env = fs.readFileSync('/home/dev/ykp/ykp-finance-v1/.env', 'utf8');
function g(k) {
  const m = env.match(new RegExp('^' + k + '=(.*)$', 'm'));
  return m ? m[1].replace(/^"|"$/g, '') : '';
}
const auth = new google.auth.JWT({
  email: g('GOOGLE_SERVICE_ACCOUNT_EMAIL'),
  key: g('GOOGLE_SERVICE_ACCOUNT_PRIVATE_KEY').replace(/\\n/g, '\n'),
  scopes: ['https://www.googleapis.com/auth/spreadsheets']
});
const s = google.sheets({ version: 'v4', auth });
const sid = g('YKP_FINANCE_SPREADSHEET_ID');
const payload = JSON.parse(fs.readFileSync('/home/dev/ykp/ykp-finance-v1/import_rev.json', 'utf8'));

const POS_DAILY_HEADERS = ['pos_id','date','brand_id','brand_name','outlet_id','outlet_name','gross_sales','net_sales','discount','refund','void','tax','service_charge','settle_cash','settle_qris','settle_card','settle_transfer','settle_marketplace','total_settlement','settlement_difference','transaction_count','aov','cashier','shift','payment_method','source','source_ref','notes','source_module','source_transaction_id','payment_source','linked_expense_id','linked_supplier_invoice_id','linked_petty_cash_id','created_by','created_at','updated_at'];
const POS_ITEMS_HEADERS = ['pos_item_id','date','brand_id','brand_name','outlet_id','outlet_name','item_name','sku','category','qty','gross_sales','discount','refund','net_sales','source','created_at'];

function toValues(headers, rows) {
  return rows.map(r => headers.map(h => r[h] ?? ''));
}

async function main() {
  // pos_daily
  const pd = toValues(POS_DAILY_HEADERS, payload.pos_daily);
  await s.spreadsheets.values.append({
    spreadsheetId: sid,
    range: 'fin_pos_daily!A1',
    valueInputOption: 'RAW',
    insertDataOption: 'INSERT_ROWS',
    requestBody: { values: pd }
  });
  console.log('pos_daily appended:', pd.length);

  // pos_items
  const pi = toValues(POS_ITEMS_HEADERS, payload.pos_items);
  await s.spreadsheets.values.append({
    spreadsheetId: sid,
    range: 'fin_pos_items!A1',
    valueInputOption: 'RAW',
    insertDataOption: 'INSERT_ROWS',
    requestBody: { values: pi }
  });
  console.log('pos_items appended:', pi.length);
}
main().then(() => console.log('DONE')).catch(e => { console.log('ERR', e.message); process.exit(1); });
"""

with sftp.file('/home/dev/ykp/ykp-finance-v1/import_rev.js', 'w') as f:
    f.write(NODE)
sftp.close()

cmd = "export PATH=/home/dev/.nvm/versions/node/v20.20.2/bin:/usr/bin:/bin:$PATH; cd /home/dev/ykp/ykp-finance-v1 && node import_rev.js"
stdin, stdout, stderr = ssh.exec_command(cmd, timeout=120)
print('=== import ===')
print(stdout.read().decode())
print(stderr.read().decode())

ssh.close()
